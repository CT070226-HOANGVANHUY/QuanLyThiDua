"""Xuất các bảng nghiệp vụ ra Excel, giữ nguyên số và ô chưa có dữ liệu."""
from __future__ import annotations

from io import BytesIO
import re

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _text(cell, value):
    """Nội dung người dùng là văn bản, không phải công thức Excel."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def workbook_bytes(tables):
    wb = Workbook()
    wb.remove(wb.active)
    for table in tables:
        title = table["title"]
        sheet_name = re.sub(r"[\\/*?:\[\]]", "-", title)[:31] or "Bảng"
        ws = wb.create_sheet(sheet_name)
        cols = table["columns"]
        width = max(1, len(cols))
        _text(ws.cell(1, 1), title)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        ws.cell(1, 1).font = Font(size=15, bold=True, color="17365D")
        notes = [table.get("source", ""), *table.get("notes", [])]
        notes = [note for note in notes if note]
        for row, note in enumerate(notes, 2):
            _text(ws.cell(row, 1), note)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
            ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 32
        header_row = len(notes) + 2
        for col, (key, label) in enumerate(cols, 1):
            cell = ws.cell(header_row, col)
            _text(cell, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = min(35, max(13, len(label) + 2))
        ws.row_dimensions[header_row].height = 40
        for row_idx, record in enumerate(table["rows"], header_row + 1):
            for col, (key, _) in enumerate(cols, 1):
                cell = ws.cell(row_idx, col)
                _text(cell, record.get(key))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.##########"
                if (row_idx - header_row) % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EDF3F8")
        ws.freeze_panes = f"C{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(width)}{max(header_row, ws.max_row)}"
        ws.print_title_rows = f"1:{header_row}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3 if width > 15 else ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    if not wb.sheetnames:
        raise ValueError("Không có bảng để xuất")
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def send_tables(tables, filename):
    return send_file(
        workbook_bytes(tables),
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
